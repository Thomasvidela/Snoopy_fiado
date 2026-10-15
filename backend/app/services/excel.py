import io
from datetime import datetime
from decimal import Decimal
from typing import Optional, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models import Cliente, Movimiento, TipoMovimiento, ahora_argentina
from app.services.resumen import parse_periodo, MESES_ES


def generar_excel_cliente(
    cliente: Cliente,
    movimientos: List[Movimiento],
    periodo: Optional[str] = None,
    saldo_anterior: Optional[Decimal] = None
) -> io.BytesIO:
    """
    Genera un archivo Excel (.xlsx) profesional con el detalle contable
    y libro mayor de movimientos del cliente.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Movimientos"
    ws.views.sheetView[0].showGridLines = True

    # Paleta de colores
    VERDE_OSCURO = "166534"
    VERDE_CLARO = "DCFCE7"
    GRIS_ENCABEZADO = "1E293B"
    GRIS_FILA_PAR = "F8FAFC"
    ROJO_SUAVE = "991B1B"
    VERDE_PAGO = "15803D"
    AMBER_CONTRA = "B45309"
    GRIS_ANULADO = "94A3B8"

    # Fuentes
    fuente_titulo = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    fuente_subtitulo = Font(name="Segoe UI", size=10, italic=True, color="FFFFFF")
    fuente_seccion = Font(name="Segoe UI", size=11, bold=True, color="1E293B")
    fuente_encabezado_tabla = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    fuente_datos = Font(name="Segoe UI", size=10)
    fuente_datos_bold = Font(name="Segoe UI", size=10, bold=True)
    fuente_anulada = Font(name="Segoe UI", size=10, strike=True, color=GRIS_ANULADO)

    # Rellenos
    fill_titulo = PatternFill(start_color=VERDE_OSCURO, end_color=VERDE_OSCURO, fill_type="solid")
    fill_encabezado_tabla = PatternFill(start_color=GRIS_ENCABEZADO, end_color=GRIS_ENCABEZADO, fill_type="solid")
    fill_kpi = PatternFill(start_color=VERDE_CLARO, end_color=VERDE_CLARO, fill_type="solid")
    fill_par = PatternFill(start_color=GRIS_FILA_PAR, end_color=GRIS_FILA_PAR, fill_type="solid")

    # Bordes
    borde_fino = Side(style="thin", color="CBD5E1")
    borde_celda = Border(left=borde_fino, right=borde_fino, top=borde_fino, bottom=borde_fino)
    borde_doble_bottom = Border(bottom=Side(style="double", color="1E293B"), top=borde_fino)

    # Alineaciones
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # 1. ENCABEZADO COMERCIO (Filas 1 y 2)
    ws.merge_cells("A1:H1")
    celda_t1 = ws["A1"]
    celda_t1.value = "DESPENSA SNOOPY - LIBRETA DE FIADOS"
    celda_t1.font = fuente_titulo
    celda_t1.fill = fill_titulo
    celda_t1.alignment = align_center

    ws.merge_cells("A2:H2")
    celda_t2 = ws["A2"]
    celda_t2.value = "Estado de Cuenta y Libro Mayor Contable"
    celda_t2.font = fuente_subtitulo
    celda_t2.fill = fill_titulo
    celda_t2.alignment = align_center

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 18

    # 2. DATOS DEL CLIENTE (Filas 4 a 6)
    alias_str = f' ("{cliente.apodo}")' if cliente.apodo else ""
    ws["A4"] = "Cliente:"
    ws["A4"].font = fuente_seccion
    ws["B4"] = f"{cliente.nombre}{alias_str}"
    ws["B4"].font = fuente_datos_bold

    ws["A5"] = "Teléfono:"
    ws["A5"].font = fuente_seccion
    ws["B5"] = cliente.telefono or "Sin teléfono"
    ws["B5"].font = fuente_datos

    ws["A6"] = "Emisión:"
    ws["A6"].font = fuente_seccion
    ws["B6"] = ahora_argentina().strftime("%d/%m/%Y %H:%M")
    ws["B6"].font = fuente_datos

    # Tarjeta de Saldo en parte derecha
    ws.merge_cells("G4:H4")
    ws["G4"] = "SALDO ACTUAL A LA FECHA"
    ws["G4"].font = Font(name="Segoe UI", size=9, bold=True, color="166534")
    ws["G4"].alignment = align_center
    ws["G4"].fill = fill_kpi

    ws.merge_cells("G5:H6")
    celda_saldo = ws["G5"]
    celda_saldo.value = float(cliente.saldo_actual)
    celda_saldo.font = Font(name="Segoe UI", size=16, bold=True, color="991B1B" if cliente.saldo_actual > 0 else "15803D")
    celda_saldo.number_format = '$ #,##0.00'
    celda_saldo.alignment = align_center
    celda_saldo.fill = fill_kpi

    for r in range(4, 7):
        for c in ["G", "H"]:
            ws[f"{c}{r}"].border = borde_celda

    # Período indicado
    if periodo:
        try:
            _, _, a, m = parse_periodo(periodo)
            nombre_mes = MESES_ES.get(m, str(m))
            txt_periodo = f"{nombre_mes} {a}"
        except Exception:
            txt_periodo = periodo
    else:
        txt_periodo = "Historial Completo"

    ws["D4"] = "Período:"
    ws["D4"].font = fuente_seccion
    ws["E4"] = txt_periodo
    ws["E4"].font = fuente_datos_bold

    fila_inicio_tabla = 8

    # 3. ENCABEZADOS DE TABLA
    encabezados = [
        ("ID", 8, align_center),
        ("Fecha / Hora", 18, align_center),
        ("Operador", 14, align_left),
        ("Tipo de Operación", 24, align_left),
        ("Detalle / Nota", 32, align_left),
        ("Monto (ARS)", 16, align_right),
        ("Saldo Resultante", 18, align_right),
        ("Estado", 14, align_center),
    ]

    for col_idx, (titulo, ancho, alineacion) in enumerate(encabezados, start=1):
        cell = ws.cell(row=fila_inicio_tabla, column=col_idx, value=titulo)
        cell.font = fuente_encabezado_tabla
        cell.fill = fill_encabezado_tabla
        cell.alignment = alineacion
        cell.border = borde_celda

    ws.row_dimensions[fila_inicio_tabla].height = 24

    # 4. FILAS DE MOVIMIENTOS
    fila_actual = fila_inicio_tabla + 1

    # Si hay saldo anterior especificado, insertar fila informativa
    if saldo_anterior is not None:
        ws.cell(row=fila_actual, column=1, value="-").alignment = align_center
        ws.cell(row=fila_actual, column=2, value="Inicio de Período").alignment = align_center
        ws.cell(row=fila_actual, column=3, value="Sistema").alignment = align_left
        ws.cell(row=fila_actual, column=4, value="Saldo Anterior").alignment = align_left
        ws.cell(row=fila_actual, column=5, value="Balance arrastrado del período previo").alignment = align_left
        ws.cell(row=fila_actual, column=6, value="-").alignment = align_right
        
        celda_s_ant = ws.cell(row=fila_actual, column=7, value=float(saldo_anterior))
        celda_s_ant.number_format = '$ #,##0.00'
        celda_s_ant.alignment = align_right
        celda_s_ant.font = fuente_datos_bold

        ws.cell(row=fila_actual, column=8, value="Inicial").alignment = align_center

        for c_idx in range(1, 9):
            ws.cell(row=fila_actual, column=c_idx).border = borde_celda
            ws.cell(row=fila_actual, column=c_idx).fill = PatternFill(start_color="F1F5F9", fill_type="solid")

        fila_actual += 1

    total_cargos_periodo = Decimal("0.00")
    total_pagos_periodo = Decimal("0.00")

    for mov in movimientos:
        es_anulado = mov.anulado
        tipo = mov.tipo
        fuente_fila = fuente_anulada if es_anulado else fuente_datos

        # Determinar texto legible del tipo
        if tipo == TipoMovimiento.CARGO:
            tipo_txt = "Compra Fiada (Cargo)"
            if not es_anulado:
                total_cargos_periodo += mov.monto
        elif tipo == TipoMovimiento.PAGO:
            tipo_txt = "Pago Recibido"
            if not es_anulado:
                total_pagos_periodo += mov.monto
        elif tipo == TipoMovimiento.ANULACION_CARGO:
            tipo_txt = "Contra-asiento (Anulación Cargo)"
        elif tipo == TipoMovimiento.ANULACION_PAGO:
            tipo_txt = "Contra-asiento (Anulación Pago)"
        else:
            tipo_txt = str(tipo)

        operador_nombre = mov.operador.nombre if mov.operador else "N/A"
        estado_txt = "ANULADO" if es_anulado else "Activo"

        # Celda 1: ID
        c1 = ws.cell(row=fila_actual, column=1, value=mov.id)
        c1.alignment = align_center

        # Celda 2: Fecha y Hora
        c2 = ws.cell(row=fila_actual, column=2, value=mov.fecha_creacion.strftime("%d/%m/%Y %H:%M"))
        c2.alignment = align_center

        # Celda 3: Operador
        c3 = ws.cell(row=fila_actual, column=3, value=operador_nombre)
        c3.alignment = align_left

        # Celda 4: Tipo de Operación
        c4 = ws.cell(row=fila_actual, column=4, value=tipo_txt)
        c4.alignment = align_left

        # Celda 5: Detalle / Nota
        c5 = ws.cell(row=fila_actual, column=5, value=mov.nota or "-")
        c5.alignment = align_left

        # Celda 6: Monto (ARS)
        signo = "-" if tipo == TipoMovimiento.PAGO or tipo == TipoMovimiento.ANULACION_CARGO else "+"
        c6 = ws.cell(row=fila_actual, column=6, value=float(mov.monto) * (-1 if signo == "-" else 1))
        c6.number_format = '$ #,##0.00'
        c6.alignment = align_right

        # Color del monto si no está anulado
        if not es_anulado:
            if tipo == TipoMovimiento.CARGO:
                c6.font = Font(name="Segoe UI", size=10, bold=True, color=ROJO_SUAVE)
            elif tipo == TipoMovimiento.PAGO:
                c6.font = Font(name="Segoe UI", size=10, bold=True, color=VERDE_PAGO)
            else:
                c6.font = Font(name="Segoe UI", size=10, bold=True, color=AMBER_CONTRA)

        # Celda 7: Saldo Resultante
        c7 = ws.cell(row=fila_actual, column=7, value=float(mov.saldo_resultante))
        c7.number_format = '$ #,##0.00'
        c7.alignment = align_right
        if not es_anulado:
            c7.font = fuente_datos_bold

        # Celda 8: Estado
        c8 = ws.cell(row=fila_actual, column=8, value=estado_txt)
        c8.alignment = align_center
        if es_anulado:
            c8.font = Font(name="Segoe UI", size=9, bold=True, color="EF4444")

        # Aplicar fuente general y bordes
        for col_i in range(1, 9):
            cell = ws.cell(row=fila_actual, column=col_i)
            cell.border = borde_celda
            if es_anulado and col_i != 8:
                cell.font = fuente_anulada
            elif not es_anulado and col_i not in (6, 7):
                cell.font = fuente_fila

            # Alternar color suave de filas
            if fila_actual % 2 == 0 and not es_anulado:
                cell.fill = fill_par

        ws.row_dimensions[fila_actual].height = 20
        fila_actual += 1

    # 5. FILA DE TOTALES / RESUMEN
    ws.merge_cells(start_row=fila_actual, start_column=1, end_row=fila_actual, end_column=5)
    celda_totales = ws.cell(row=fila_actual, column=1, value="TOTALES DEL LISTADO (Cargos - Pagos)")
    celda_totales.font = fuente_datos_bold
    celda_totales.alignment = align_right

    neto_periodo = total_cargos_periodo - total_pagos_periodo
    celda_tot_monto = ws.cell(row=fila_actual, column=6, value=float(neto_periodo))
    celda_tot_monto.number_format = '$ #,##0.00'
    celda_tot_monto.font = fuente_datos_bold
    celda_tot_monto.alignment = align_right

    celda_tot_saldo = ws.cell(row=fila_actual, column=7, value=float(cliente.saldo_actual))
    celda_tot_saldo.number_format = '$ #,##0.00'
    celda_tot_saldo.font = Font(name="Segoe UI", size=11, bold=True, color="991B1B" if cliente.saldo_actual > 0 else "15803D")
    celda_tot_saldo.alignment = align_right

    ws.cell(row=fila_actual, column=8, value="").alignment = align_center

    for c_i in range(1, 9):
        ws.cell(row=fila_actual, column=c_i).border = borde_doble_bottom

    ws.row_dimensions[fila_actual].height = 24

    # Ajustar anchos de columnas
    for col_idx, (_, ancho, _) in enumerate(encabezados, start=1):
        col_letra = get_column_letter(col_idx)
        ws.column_dimensions[col_letra].width = ancho

    # Inmovilizar paneles para que la cabecera quede fija al hacer scroll
    ws.freeze_panes = f"A{fila_inicio_tabla + 1}"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
